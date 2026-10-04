"""Writes notebooks/hmha_by_hand.xlsx: the whole project on a tiny version of the task, in
L-shaped spreadsheet matrices (left matrix bottom-left, top matrix top-right, product where they
meet). Every product cell is a formula, so changing a yellow input updates everything.

Tiny task: N = 4 memory slots, R = 2 distances (1 and 3), one number per item.

  python experiments/build_by_hand_xlsx.py
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as COL

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
OUT = os.path.join(ROOT, "notebooks", "hmha_by_hand.xlsx")
N, R, OFFS = 4, 2, (1, 3)
ITEMS = [2, 1, -3, 6]             # chosen so that sum_p c(p+1) * c(p+3) = 0 (see sheet 5)
P = [2, 0, 3, 1]                  # the position each of the 4 query tokens asks about

FONT = "Arial"
FILL_IN = PatternFill("solid", fgColor="FFF2CC")    # yellow: an input you may change
FILL_A = PatternFill("solid", fgColor="DDEBF7")     # blue: left matrix
FILL_B = PatternFill("solid", fgColor="FCE4D6")     # orange: top matrix
FILL_C = PatternFill("solid", fgColor="E2EFDA")     # green: the product
FILL_X = PatternFill("solid", fgColor="EDEDED")     # grey: other computed cells
EDGE = Side(style="thin", color="BFBFBF")
BOX = Border(EDGE, EDGE, EDGE, EDGE)
TASK = "1 Task"


def ref(r, c, sheet=None, fixed=False):
    a = f"${COL(c)}${r}" if fixed else f"{COL(c)}{r}"
    return f"'{sheet}'!{a}" if sheet else a


def put(ws, r, c, v, fill=None, inp=False, fmt="General"):
    cell = ws.cell(r, c, v)
    cell.font = Font(name=FONT, size=10, color="0000FF" if inp else "000000")
    if fill is not None:
        cell.fill, cell.border = fill, BOX
    cell.number_format = fmt
    cell.alignment = Alignment(horizontal="center", vertical="center")
    return cell


def text(ws, r, c, s, bold=False, size=10, color="000000", italic=False):
    cell = ws.cell(r, c, s)
    cell.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def label(ws, r, c, s):
    cell = ws.cell(r, c, s)
    cell.font = Font(name=FONT, size=9, color="595959", italic=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def matrix(ws, r, c, vals, fill, inp=False, fmt="General"):
    for i, row in enumerate(vals):
        for j, v in enumerate(row):
            put(ws, r + i, c + j, v, fill, inp=inp and not (isinstance(v, str) and v.startswith("=")), fmt=fmt)
    return r, c, len(vals), len(vals[0])


def lblock(ws, top, left, A, B, nameA, nameB, nameC, rowlab=None, collab=None,
           inpA=False, inpB=False, fmtA="General", fmtB="General", fmtC="General"):
    """C = A B laid out as an L: B top-right, A bottom-left, C where they meet.
    A, B: 2-D lists of numbers or formulas. Returns the top-left cell and size of A, B and C."""
    n, k, m = len(A), len(B), len(B[0])
    rb, cb = top + 2, left + k + 1                    # B: below its label and column headers
    ra, ca = rb + k, left                             # A: directly below-left of B
    text(ws, top, cb, nameB, bold=True)
    text(ws, ra - 1, ca, nameA, bold=True)
    if collab:
        for j, s in enumerate(collab):
            label(ws, top + 1, cb + j, s)
    if rowlab:
        for i, s in enumerate(rowlab):
            label(ws, ra + i, ca - 1, s)
    a = matrix(ws, ra, ca, A, FILL_A, inp=inpA, fmt=fmtA)
    b = matrix(ws, rb, cb, B, FILL_B, inp=inpB, fmt=fmtB)
    prod = [["=" + "+".join(f"{ref(ra + i, ca + t)}*{ref(rb + t, cb + j)}" for t in range(k))
             for j in range(m)] for i in range(n)]
    c = matrix(ws, ra, cb, prod, FILL_C, fmt=fmtC)
    text(ws, ra + n, cb, nameC, bold=True)
    return a, b, c


def cells(block, sheet=None):
    r, c, n, m = block
    return [[ref(r + i, c + j, sheet) for j in range(m)] for i in range(n)]


def widths(ws, first=2.5, rest=8.5, n=40):
    ws.column_dimensions["A"].width = first
    for c in range(2, n):
        ws.column_dimensions[COL(c)].width = rest


wb = Workbook()

# ===================================================================== 0 Read me
ws = wb.active
ws.title = "0 Read me"
widths(ws)
text(ws, 1, 2, "The collateral rule by hand: the whole project on a tiny task", bold=True, size=14)
lines = [
    "Each sheet is one step. Work through them in order: 1 Task, 2 One head, 3 Two heads, 4 Why R heads,",
    "5 Collateral value, 6 The rule, 7 Backup heads.",
    "",
    "Colours:  yellow = an input you may change (blue numbers)   blue = left matrix   orange = top matrix",
    "          green = the product (a formula)   grey = other computed cells",
    "",
    "How to read an L-shape. To multiply A times B, put B above and to the right, A below and to the left.",
    "Each green cell sits in A's row and B's column: multiply them entry by entry and add. Try it below,",
    "then click any green cell to see the formula Excel uses: it is exactly your hand calculation.",
]
for i, s in enumerate(lines):
    text(ws, 3 + i, 2, s)
lblock(ws, 13, 3, [[1, 2], [3, 4]], [[5, 6], [7, 8]], "A", "B", "A B",
       rowlab=["row 1", "row 2"], collab=["col 1", "col 2"], inpA=True, inpB=True, fmtC="General")
text(ws, 21, 2, "By hand: row 1 times col 1 = 1*5 + 2*7 = 19.   row 2 times col 2 = 3*6 + 4*8 = 50.")
text(ws, 23, 2, "The tiny task: 4 memory slots, 2 distances (1 and 3), one number per item. The real experiments use")
text(ws, 24, 2, "16 slots, 4 distances (1, 5, 9, 13), 16 numbers per item and 32 heads: same steps, bigger matrices.")

# ===================================================================== 1 Task
ws = wb.create_sheet(TASK)
widths(ws, rest=12)
text(ws, 1, 2, "Step 1. The task", bold=True, size=14)
text(ws, 2, 2, "A memory of 4 slots holds one item each. A query names a position p and must return the items")
text(ws, 3, 2, "1 and 3 slots ahead, counting round like a clock (mod 4). Change the yellow cells and watch the targets.")
text(ws, 5, 2, "slot j")
text(ws, 6, 2, "item c_j")
for j in range(N):
    put(ws, 5, 3 + j, j, FILL_X)
    put(ws, 6, 3 + j, ITEMS[j], FILL_IN, inp=True)
ITEM = [ref(6, 3 + j, TASK, fixed=True) for j in range(N)]
ITEMS_RANGE = f"$C$6:${COL(2 + N)}$6"
hdr = ["query i", "p", "p + 1", "(p+1) mod 4", "target 1", "p + 3", "(p+3) mod 4", "target 2"]
for j, s in enumerate(hdr):
    text(ws, 9, 2 + j, s, bold=True)
for i in range(N):
    r = 10 + i
    put(ws, r, 2, i, FILL_X)
    put(ws, r, 3, P[i], FILL_IN, inp=True)
    for k, off in enumerate(OFFS):
        c0 = 4 + 3 * k
        put(ws, r, c0, f"=C{r}+{off}", FILL_X)
        put(ws, r, c0 + 1, f"=MOD({COL(c0)}{r},{N})", FILL_X)
        put(ws, r, c0 + 2, f"=INDEX({ITEMS_RANGE},{COL(c0 + 1)}{r}+1)", FILL_C)
PCELL = [ref(10 + i, 3, TASK, fixed=True) for i in range(N)]
TGT = [[ref(10 + i, 6, TASK, fixed=True), ref(10 + i, 9, TASK, fixed=True)] for i in range(N)]
text(ws, 14, 2, "target 1 = the item in slot (p+1) mod 4;  target 2 = the item in slot (p+3) mod 4.", italic=True)
text(ws, 15, 2, "By hand: query 0 has p = 2. (2+1) mod 4 = 3, so target 1 = c_3 = 6. (2+3) mod 4 = 1, so target 2 = c_1 = 1.")
text(ws, 16, 2, "Why mod: without it p = 2 would need slot 5, which does not exist. Wrapping round gives every query")
text(ws, 17, 2, "both answers inside the memory, and 'look 1 ahead' is the same rule at every position.")
text(ws, 19, 2, "As matrices (used on the next sheets): memory row j = [one-hot of j | c_j], query row i = one-hot of p_i,")
text(ws, 20, 2, "target row i = [c_(p+1) | c_(p+3)]. The real code adds a little noise to every vector.")
text(ws, 22, 2, "What are 'memory' and 'slots'? Yes, the memory is just a list: [2, 1, -3, 6]. A slot is a position", bold=True)
text(ws, 23, 2, "in that list (0, 1, 2, 3), like an index. In plain Python the whole task is:")
text(ws, 24, 3, "memory = [2, 1, -3, 6];   answer(p) = (memory[(p + 1) % 4], memory[(p + 3) % 4])")
text(ws, 25, 2, "So why the one-hot labels? Attention cannot see list positions: it compares the query with every row and")
text(ws, 26, 2, "treats the rows like an unordered pile. So each memory row carries its own position as a code next to its")
text(ws, 27, 2, "item (slot 2 = [0, 0, 1, 0 | -3]), the query carries a code for p, and the head learns to match codes.")
text(ws, 28, 2, "The items change every example, so the model cannot memorise them; it can only learn WHERE to look.")


# ===================================================================== a head, built by hand
def build_head(ws, top, shift, tag):
    """One attention head that looks `shift` slots ahead. Returns the cells of its output O (4 x 1)
    and the next free row."""
    slots = [f"slot {j}" for j in range(N)]
    queries = [f"p = {P[i]}" for i in range(N)]
    # memory Y: [one-hot of j | item]
    Y = [[1 if t == j else 0 for t in range(N)] + [f"={ITEM[j]}"] for j in range(N)]
    ycol = [f"lab {t}" for t in range(N)] + ["item"]
    # keys: K = Y W_K, W_K copies the slot label and ignores the item
    WK = [[1 if (t == u) else 0 for u in range(N)] for t in range(N)] + [[0] * N]
    text(ws, top, 2, f"(a) Keys: K = Y W_K. W_K keeps each slot's label and ignores its item, so key j = one-hot of j.",
         italic=True)
    _, _, K = lblock(ws, top + 1, 3, Y, WK, "Y (memory)", "W_K", "K = Y W_K",
                     rowlab=slots, collab=[f"k{u}" for u in range(N)], inpB=True, fmtC="General")
    top = K[0] + N + 2
    # values: V = Y W_V, W_V keeps only the item
    WV = [[0]] * N + [[1]]
    text(ws, top, 2, "(b) Values: V = Y W_V. W_V keeps only the item, so value j = c_j.", italic=True)
    _, _, V = lblock(ws, top + 1, 3, Y, WV, "Y (memory)", "W_V", "V = Y W_V",
                     rowlab=slots, collab=["v"], inpB=True, fmtC="General")
    top = V[0] + N + 2
    # queries: Q = X W_Q, W_Q shifts the label by `shift`
    X = [[f"=IF({PCELL[i]}={t},1,0)" for t in range(N)] for i in range(N)]
    WQ = [[1 if u == (t + shift) % N else 0 for u in range(N)] for t in range(N)]
    text(ws, top, 2, f"(c) Queries: Q = X W_Q. Row t of W_Q has its 1 in column (t + {shift}) mod 4, so the query for p "
                     f"becomes one-hot of p + {shift}: this head looks {shift} ahead.", italic=True)
    _, _, Q = lblock(ws, top + 1, 3, X, WQ, "X (queries, one-hot of p)", f"W_Q (shift by {shift})", "Q = X W_Q",
                     rowlab=queries, collab=[f"q{u}" for u in range(N)], inpB=True, fmtC="General")
    top = Q[0] + N + 2
    # scores: S = Q K^T
    text(ws, top, 2, "(d) Scores: S = Q K^T. Row i, column j = query i dotted with key j: 1 where j = p + shift.",
         italic=True)
    Qc, Kc = cells(Q), cells(K)
    KT = [[f"={Kc[j][u]}" for j in range(N)] for u in range(N)]
    _, _, S = lblock(ws, top + 1, 3, [[f"={c}" for c in row] for row in Qc], KT, "Q (copied)", "K^T (K turned on its side)",
                     "S = Q K^T", rowlab=queries, collab=slots, fmtC="General")
    top = S[0] + N + 2
    # softmax, row by row
    text(ws, top, 2, "(e) Attention = softmax of (sharpness x S), row by row: e^(sharpness x score) / sum of the row's.",
         italic=True)
    text(ws, top + 1, 3, "sharpness")
    put(ws, top + 1, 4, 10, FILL_IN, inp=True)
    sharp = ref(top + 1, 4, fixed=True)
    Sc = cells(S)
    att = [[f"=EXP({sharp}*{Sc[i][j]})/(" + "+".join(f"EXP({sharp}*{Sc[i][t]})" for t in range(N)) + ")"
            for j in range(N)] for i in range(N)]
    text(ws, top + 2, 3, "attention", bold=True)
    for j in range(N):
        label(ws, top + 2, 4 + j, f"slot {j}")
    A = matrix(ws, top + 3, 4, att, FILL_X, fmt="0.000")
    for i in range(N):
        label(ws, top + 3 + i, 3, queries[i])
        put(ws, top + 3 + i, 4 + N, f"=SUM({COL(4)}{top + 3 + i}:{COL(3 + N)}{top + 3 + i})", FILL_X, fmt="0.000")
    label(ws, top + 2, 4 + N, "row sum")
    top = top + 3 + N + 1
    # output: O = attention V
    text(ws, top, 2, "(f) Output: O = attention x V, a weighted average of the items. Compare with the target.",
         italic=True)
    Ac, Vc = cells(A), cells(V)
    _, _, O = lblock(ws, top + 1, 3, [[f"={c}" for c in row] for row in Ac], [[f"={Vc[j][0]}"] for j in range(N)],
                     "attention (copied)", "V (copied)", f"O_{tag}", rowlab=queries, collab=["out"], fmtA="0.000", fmtC="0.000")
    tcol = O[1] + 2
    text(ws, O[0] - 1, tcol, f"target c(p+{shift})", bold=True)
    k = OFFS.index(shift)
    for i in range(N):
        put(ws, O[0] + i, tcol, f"={TGT[i][k]}", FILL_X)
    return cells(O), O[0] + N + 2


# ===================================================================== 2 One head
ws = wb.create_sheet("2 One head")
widths(ws)
text(ws, 1, 2, "Step 2. One attention head, built by hand (it looks 1 slot ahead)", bold=True, size=14)
text(ws, 2, 2, "Follow (a) to (f). The head scores every slot, turns the scores into weights that add up to 1, and")
text(ws, 3, 2, "returns the weighted average of the items. Try: change W_Q so the 1s move, or lower the sharpness to 1.")
O1, _ = build_head(ws, 5, 1, "1")

# ===================================================================== 3 Two heads and valves
ws = wb.create_sheet("3 Two heads")
widths(ws)
text(ws, 1, 2, "Step 3. A second head (looks 3 ahead), then both heads through their valves", bold=True, size=14)
text(ws, 2, 2, "Head 2 is built exactly like head 1; only W_Q differs (shift by 3). Its output is at the bottom of (f).")
O2, top = build_head(ws, 4, 3, "2")
text(ws, top, 2, "(g) Valves and read-out: prediction = [g_1 O_1 | g_2 O_2] W_O. This is the line out * gate in the code.",
     italic=True)
text(ws, top + 1, 3, "valve g_1")
put(ws, top + 1, 4, 1, FILL_IN, inp=True)
text(ws, top + 1, 6, "valve g_2")
put(ws, top + 1, 7, 1, FILL_IN, inp=True)
g1, g2 = ref(top + 1, 4, fixed=True), ref(top + 1, 7, fixed=True)
O1_other = [[f"'2 One head'!{c[0]}"] for c in O1]
Z = [[f"={g1}*{O1_other[i][0]}", f"={g2}*{O2[i][0]}"] for i in range(N)]
_, _, PRED = lblock(ws, top + 3, 3, Z, [[1, 0], [0, 1]], "[g_1 O_1 | g_2 O_2]", "W_O", "prediction",
                    rowlab=[f"p = {p}" for p in P], collab=["ans 1", "ans 2"], inpB=True, fmtA="0.000", fmtC="0.000")
pr = cells(PRED)
r0, c0 = PRED[0], PRED[1] + 3
text(ws, r0 - 1, c0, "target", bold=True)
text(ws, r0 - 1, c0 + 3, "squared error", bold=True)
for i in range(N):
    for k in range(R):
        put(ws, r0 + i, c0 + k, f"={TGT[i][k]}", FILL_X)
        put(ws, r0 + i, c0 + 3 + k, f"=({pr[i][k]}-{ref(r0 + i, c0 + k)})^2", FILL_X, fmt="0.000")
text(ws, r0 + N + 1, c0 + 3, "loss (average)", bold=True)
put(ws, r0 + N + 1, c0 + 5, f"=AVERAGE({COL(c0 + 3)}{r0}:{COL(c0 + 4)}{r0 + N - 1})", FILL_C, fmt="0.000")
text(ws, r0 + N + 3, 2, "Try: set valve g_2 to 0. Column 2 of the prediction becomes 0 whatever head 2 computes, and the loss")
text(ws, r0 + N + 4, 2, "jumps. A head with its valve at 0 is removed exactly: it adds nothing, so training sends it no gradient.")

# ===================================================================== 4 Why R heads
ws = wb.create_sheet("4 Why R heads")
widths(ws, rest=10)
text(ws, 1, 2, "Step 4. One head, one equation: why the task needs exactly R heads", bold=True, size=14)
text(ws, 2, 2, "A head returns ONE weighted average of the items: one equation in the unknown items. Take query 0's two")
text(ws, 3, 2, "unknowns a = c(p+1) and b = c(p+3). Each row of the mixing matrix is one head's recipe.")
_, _, U = lblock(ws, 5, 3, [[1, 0], [0, 1]], [[f"={TGT[0][0]}"], [f"={TGT[0][1]}"]], "mixing (one row per head)",
                 "unknowns (a; b)", "what the heads hand back", rowlab=["head 1", "head 2"], collab=["value"],
                 inpA=True, fmtC="General")
r = U[0]
det = f"({ref(r, 3)}*{ref(r + 1, 4)}-{ref(r, 4)}*{ref(r + 1, 3)})"
text(ws, r + 4, 2, "Can we get a and b back from the two values? Only if the two recipes are different:")
text(ws, r + 5, 3, "determinant = m11 m22 - m12 m21")
put(ws, r + 5, 7, f"={det}", FILL_X, fmt="General")
text(ws, r + 6, 3, "recovered a")
put(ws, r + 6, 7, f"=IF({det}=0,\"cannot\",({ref(r + 1, 4)}*{ref(r, 6)}-{ref(r, 4)}*{ref(r + 1, 6)})/{det})", FILL_C, fmt="General")
text(ws, r + 7, 3, "recovered b")
put(ws, r + 7, 7, f"=IF({det}=0,\"cannot\",({ref(r, 3)}*{ref(r + 1, 6)}-{ref(r + 1, 3)}*{ref(r, 6)})/{det})", FILL_C, fmt="General")
text(ws, r + 9, 2, "Try: (1) mixing rows (0.5, 0.5) and (0.5, -0.5): still recovered. (2) make head 2 a copy of head 1:")
text(ws, r + 10, 2, "determinant 0, 'cannot'. Two copies give the same equation twice, so they count once.")
text(ws, r + 12, 2, "So with k different heads you can recover k of the R items; the rest stay unknown, and the best loss is")
text(ws, r + 13, 2, "(1 - k/R) x the loss of a model that learned nothing. Checked on a trained model (R = 4, real task):")
for j, s in enumerate(["different heads k", "measured loss", "predicted (1 - k/4) x 1.008"]):
    text(ws, r + 15, 3 + 2 * j, s, bold=True)
for i, (k, meas) in enumerate([(4, 0.000), (3, 0.254), (2, 0.507), (1, 0.758)]):
    put(ws, r + 16 + i, 3, k, FILL_X)
    put(ws, r + 16 + i, 5, meas, FILL_X, fmt="0.000")
    put(ws, r + 16 + i, 7, f"=(1-{ref(r + 16 + i, 3)}/4)*1.008", FILL_C, fmt="0.000")
text(ws, r + 21, 2, "Source of the measured losses: results/proposal/equations.pkl (experiments/equations_test.py).",
     italic=True, color="595959")

# ===================================================================== 5 Collateral value
ws = wb.create_sheet("5 Collateral value")
widths(ws)
text(ws, 1, 2, "Step 5. The collateral value: what nobody else can cover", bold=True, size=14)
text(ws, 2, 2, "Three trained heads. Head A looks 1 ahead, head B looks 3 ahead, head C is an exact COPY of A. Their")
text(ws, 3, 2, "outputs are the items they fetch (as in steps 2 and 3). The read-out W_O turns them into the two answers.")
text(ws, 4, 2, "A head's collateral value = how much the loss rises when it is removed AND the others re-fit W_O.")
ZA = [f"={TGT[i][0]}" for i in range(N)]
ZB = [f"={TGT[i][1]}" for i in range(N)]
cases = [
    ("Case 1. All heads on.", (1, 1, 1), [[0.5, 0], [0, 1], [0.5, 0]]),
    ("Case 2. A off, W_O unchanged (ordinary importance: nobody re-fits).", (0, 1, 1), [[0.5, 0], [0, 1], [0.5, 0]]),
    ("Case 3. A off, W_O re-fitted by hand: C takes A's weight (collateral value of A).", (0, 1, 1), [[0, 0], [0, 1], [1, 0]]),
    ("Case 4. B off, best re-fit: nobody else carries c(p+3), so its column gets weight 0 (collateral value of B).",
     (1, 0, 1), [[0.5, 0], [0, 0], [0.5, 0]]),
]
LOSS = []
top = 6
for title, g, W in cases:
    text(ws, top, 2, title, bold=True)
    text(ws, top + 1, 3, "valves g_A, g_B, g_C")
    for j in range(3):
        put(ws, top + 1, 7 + j, g[j], FILL_IN, inp=True)
    gref = [ref(top + 1, 7 + j, fixed=True) for j in range(3)]
    Zg = [[f"={gref[0]}*{ZA[i][1:]}", f"={gref[1]}*{ZB[i][1:]}", f"={gref[2]}*{ZA[i][1:]}"] for i in range(N)]
    _, _, Pc = lblock(ws, top + 2, 3, Zg, W, "[g_A A | g_B B | g_C C]", "W_O", "prediction",
                      rowlab=[f"p = {p}" for p in P], collab=["ans 1", "ans 2"], inpB=True, fmtC="General")
    pr = cells(Pc)
    r0, c0 = Pc[0], Pc[1] + 3
    text(ws, r0 - 1, c0, "target", bold=True)
    text(ws, r0 - 1, c0 + 3, "squared error", bold=True)
    for i in range(N):
        for k in range(R):
            put(ws, r0 + i, c0 + k, f"={TGT[i][k]}", FILL_X)
            put(ws, r0 + i, c0 + 3 + k, f"=({pr[i][k]}-{ref(r0 + i, c0 + k)})^2", FILL_X, fmt="General")
    text(ws, r0 + N, c0 + 3, "loss", bold=True)
    put(ws, r0 + N, c0 + 4, f"=AVERAGE({COL(c0 + 3)}{r0}:{COL(c0 + 4)}{r0 + N - 1})", FILL_C, fmt="General")
    LOSS.append(ref(r0 + N, c0 + 4, "5 Collateral value", fixed=True))
    top = r0 + N + 3
text(ws, top, 2, "Check of case 4: the best weight for answer 2 on head A's output is sum(A x B) / sum(A x A):", italic=True)
rngA = f"'{TASK}'!$F$10:$F${9 + N}"
rngB = f"'{TASK}'!$I$10:$I${9 + N}"
put(ws, top, 11, f"=SUMPRODUCT({rngA},{rngB})/SUMPRODUCT({rngA},{rngA})", FILL_C, fmt="General")
text(ws, top + 1, 2, "It is 0 for these items, so A (and its copy C) can cover nothing of B's job.", italic=True)
top += 3
text(ws, top, 2, "Summary", bold=True)
for c, s in ((2, "head"), (4, "ordinary importance (no re-fit)"), (8, "collateral value (re-fit)")):
    text(ws, top + 1, c, s, bold=True)
L1, L2, L3, L4 = [s.split("!")[1] for s in LOSS]
rows = [("A", f"={L2}-{L1}", f"={L3}-{L1}"), ("B", f"={L4}-{L1}", f"={L4}-{L1}"), ("C (copy of A)", f"={L2}-{L1}", f"={L3}-{L1}")]
VAL = {}
for i, (h, imp, col) in enumerate(rows):
    text(ws, top + 2 + i, 2, h)
    put(ws, top + 2 + i, 4, imp, FILL_X, fmt="General")
    put(ws, top + 2 + i, 8, col, FILL_C, fmt="General")
    VAL[h[0]] = ref(top + 2 + i, 8, "5 Collateral value", fixed=True)
text(ws, top + 6, 2, "Ordinary importance says A and C both matter (removing either alone hurts). The collateral value says")
text(ws, top + 7, 2, "each is worth 0, because its copy covers it. Only B is irreplaceable. That is why the rule removes copies.")

# ===================================================================== 6 The rule
ws = wb.create_sheet("6 The rule")
widths(ws, rest=12)
text(ws, 1, 2, "Step 6. The rule: close the cheapest head while it is worth less than the price", bold=True, size=14)
text(ws, 2, 2, "Every 25 training steps: value every open head, close the cheapest if its value is below the price,")
text(ws, 3, 2, "and fade it out slowly. Repeat until every open head is worth more than the price.")
text(ws, 5, 2, "price")
put(ws, 5, 4, 1, FILL_IN, inp=True)
price = "$D$5"
text(ws, 7, 2, "Round 1 (all three heads open)", bold=True)
for j, h in enumerate("ABC"):
    text(ws, 8, 3 + j, h, bold=True)
    put(ws, 9, 3 + j, f"={VAL[h]}", FILL_X, fmt="General")
text(ws, 9, 2, "value")
text(ws, 10, 2, "cheapest")
put(ws, 10, 3, "=MIN(C9:E9)", FILL_X, fmt="General")
put(ws, 10, 4, "=INDEX(C8:E8,MATCH(C10,C9:E9,0))", FILL_X)
text(ws, 11, 2, "decision")
put(ws, 11, 3, f"=IF(C10<{price},\"close \"&D10,\"stop\")", FILL_C)
text(ws, 13, 2, "Round 2 (A closed). Now C is the only head carrying c(p+1), so removing C would lose answer 1 entirely:",
     bold=True)
for j, h in enumerate(["A", "B", "C"]):
    text(ws, 14, 3 + j, h, bold=True)
put(ws, 15, 3, "closed", FILL_X)
put(ws, 15, 4, "=D9", FILL_X, fmt="General")
put(ws, 15, 5, f"=SUMPRODUCT({rngA},{rngA})/{2 * N}", FILL_X, fmt="General")
text(ws, 15, 2, "value")
text(ws, 16, 2, "cheapest")
put(ws, 16, 3, "=MIN(D15:E15)", FILL_X, fmt="General")
put(ws, 16, 4, "=INDEX(D14:E14,MATCH(C16,D15:E15,0))", FILL_X)
text(ws, 17, 2, "decision")
put(ws, 17, 3, f"=IF(C16<{price},\"close \"&D16,\"stop\")", FILL_C)
text(ws, 17, 4, "(stop = every open head is worth more than the price)", italic=True)
text(ws, 19, 2, "C's value in round 2 = average over both answers of c(p+1)^2 = sum of c(p+1)^2 / 8.", italic=True)
text(ws, 20, 2, "Result: B and C stay, 2 heads for a 2-distance task. Nobody told the rule the answer is 2.")
text(ws, 22, 2, "Try: raise the price above both round-2 values. The rule then closes a needed head and the task breaks.")
text(ws, 23, 2, "Lower it to 0: nothing is ever closed. In the experiments the price is 3% of a know-nothing model's loss.")
text(ws, 24, 2, "In training, a closed head fades out over 100 steps, so the other heads adjust while it goes.")

# ===================================================================== 7 Backup heads
ws = wb.create_sheet("7 Backup heads")
widths(ws, rest=12)
text(ws, 1, 2, "Step 7. Backup heads when heads can fail", bold=True, size=14)
text(ws, 2, 2, "If every head fails (valve forced to 0) with chance p at each step, a spare head is worth keeping when it")
text(ws, 3, 2, "is likely to be needed: it works AND fewer than R of the other heads work. Its value with B heads open:")
text(ws, 4, 3, "v(B) = (1 - p) / R x P[ Binomial(B - 1, 1 - p) <= R - 1 ]")
for i, (s, v) in enumerate([("R (heads needed)", 4), ("p (chance a head fails)", 0.1), ("price (fraction of a know-nothing loss)", 0.03)]):
    text(ws, 6 + i, 2, s)
    put(ws, 6 + i, 6, v, FILL_IN, inp=True)
Rr, pp, pr_ = "$F$6", "$F$7", "$F$8"
for j, s in enumerate(["B heads open", "P[fewer than R others work]", "value v(B)", "keep?"]):
    text(ws, 10, 2 + 2 * j, s, bold=True)
for i in range(12):
    r = 11 + i
    put(ws, r, 2, i + 1, FILL_X)
    put(ws, r, 4, f"=IF(B{r}-1<={Rr}-1,1,BINOMDIST({Rr}-1,B{r}-1,1-{pp},TRUE))", FILL_X, fmt="0.0000")
    put(ws, r, 6, f"=(1-{pp})/{Rr}*D{r}", FILL_C, fmt="0.0000")
    put(ws, r, 8, f"=IF(F{r}>={pr_},\"keep\",\"drop\")", FILL_X)
text(ws, 24, 2, "predicted heads kept", bold=True)
put(ws, 24, 6, f"=COUNTIF(F11:F22,\">=\"&{pr_})", FILL_C)
text(ws, 26, 2, "By hand (R = 4, p = 0.1): B = 5 needs 'fewer than 4 of the other 4 work' = 1 - 0.9^4 = 0.344,")
text(ws, 27, 2, "so v(5) = 0.9/4 x 0.344 = 0.077 > 0.03: keep. B = 6 gives 0.018 < 0.03: drop. Predicted: 5 heads.")
text(ws, 29, 2, "Measured in training (3 seeds each), written down BEFORE the runs:", bold=True)
meas = [(2, 0.05, "3, 3, 3"), (2, 0.1, "3, 3, 3"), (2, 0.2, "4, 4, 4"), (2, 0.3, "5, 4, 5"),
        (4, 0.05, "5, 5, 4"), (4, 0.1, "5, 5, 5"), (4, 0.2, "6, 6, 6"), (4, 0.3, "7, 7, 7")]
for j, s in enumerate(["R", "p", "measured heads kept"]):
    text(ws, 30, 2 + 2 * j, s, bold=True)
for i, (Rv, pv, m) in enumerate(meas):
    put(ws, 31 + i, 2, Rv, FILL_X)
    put(ws, 31 + i, 4, pv, FILL_X)
    put(ws, 31 + i, 6, m, FILL_X)
text(ws, 40, 2, "Source: results/robust (experiments/run.py with --head_dropout). Put R and p above to compare.",
     italic=True, color="595959")

# ===================================================================== side panels: real sizes and real code
import inspect, sys, textwrap
sys.path.insert(0, os.path.join(ROOT, "src"))
from hemo import tasks as h_tasks, model as h_model, train as h_train

PANEL = 16                                        # first column of the side panel
CODE_FONT = Font(name="Consolas", size=9, color="1F1F1F")
FILL_CODE = PatternFill("solid", fgColor="F7F7F7")


def cut(lines, start=None, end=None, strip_doc=True):
    """The lines from the one containing `start` to the one containing `end`, docstrings removed."""
    if start:
        lines = lines[next(i for i, l in enumerate(lines) if start in l):]
    if end:
        lines = lines[:next(i for i, l in enumerate(lines) if end in l) + 1]
    if strip_doc:
        out, inside = [], False
        for l in lines:
            s = l.strip()
            if not inside and s.startswith('"""'):
                inside = not (s.endswith('"""') and len(s) > 3)
                continue
            if inside:
                inside = not s.endswith('"""')
                continue
            out.append(l)
        lines = out
    return textwrap.dedent("\n".join(lines)).splitlines()


def src(obj, start=None, end=None):
    return cut(inspect.getsource(obj).splitlines(), start, end)


def file_src(rel, start, end):
    return cut(open(os.path.join(ROOT, rel), encoding="utf-8").read().splitlines(), start, end)


def side_panel(ws, dims, code):
    """dims: rows of (object, tiny, real, note). code: list of (title, lines)."""
    for c, w in zip(range(PANEL, PANEL + 6), (24, 16, 24, 46, 14, 14)):
        ws.column_dimensions[COL(c)].width = w
    text(ws, 5, PANEL, "Sizes: this tiny task vs the real experiments", bold=True, size=11)
    for j, s in enumerate(("object", "tiny (this sheet)", "real experiments", "note")):
        text(ws, 6, PANEL + j, s, bold=True)
    for i, row in enumerate(dims):
        for j, v in enumerate(row):
            cell = ws.cell(7 + i, PANEL + j, v)
            cell.font = Font(name=FONT, size=9, color="000000")
            cell.fill, cell.border = (FILL_X if j < 3 else PatternFill()), (BOX if j < 3 else Border())
            cell.alignment = Alignment(horizontal="left", vertical="center")
    r = 7 + len(dims) + 2
    for title, lines in code:
        text(ws, r, PANEL, title, bold=True, size=11)
        r += 1
        for line in lines:
            for c in range(PANEL, PANEL + 6):
                ws.cell(r, c).fill = FILL_CODE
            cell = ws.cell(r, PANEL, line)
            cell.font = CODE_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")
            r += 1
        r += 1


B = "B = examples per batch (128 in training, 512 per probe, 1024 for validation)"
side_panel(wb["0 Read me"], [
    ("memory slots N", "4", "16", ""),
    ("distances R (= heads needed)", "2: (1, 3)", "4: (1, 5, 9, 13)", ""),
    ("numbers per item m", "1", "16", ""),
    ("vector length d_model", "5", "128", "tiny: 4 label + 1 item; real: 16 label + 16 item + noise"),
    ("heads H", "2 or 3", "32", ""),
    ("head size d_k", "4 (1 for values)", "32", "the real code uses the same d_k for queries, keys and values"),
    ("answer length R x m", "2", "64", ""),
    ("training steps", "-", "4000", "Adam, learning rate 0.002, batch 128"),
], [])

side_panel(wb[TASK], [
    ("query tokens per example", "4", "16", "one per position p"),
    ("X (queries)", "(4, 4)", "(B, 16, 128)", B),
    ("Y (memory)", "(4, 5)", "(B, 16, 128)", "row j = [one-hot of j | item | noise]"),
    ("item c_j", "1 number", "16 numbers", "content = torch.randn(B, N, m)"),
    ("T (targets)", "(4, 2)", "(B, 16, 64)", "R items of 16 numbers, side by side"),
    ("offsets", "(1, 3)", "(1, 5, 9, 13)", "1 + r * N / R for r = 0 .. R-1"),
], [("The code: src/hemo/tasks.py", src(h_tasks.offsets_for) + [""] +
     src(h_tasks.make_batch, 'elif cfg.task == "multi_relation"', 'T = torch.cat'))])

side_panel(wb["2 One head"], [
    ("W_Q, one head", "(4, 4)", "(128, 32)", "all 32 heads stored together: W_q is 128 -> 1024"),
    ("W_K, one head", "(5, 4)", "(128, 32)", "W_k: 128 -> 1024"),
    ("W_V, one head", "(5, 1)", "(128, 32)", "W_v: 128 -> 1024"),
    ("Q = X W_Q", "(4, 4)", "(B, 32, 16, 32)", "after _split: (examples, heads, tokens, d_k)"),
    ("K = Y W_K", "(4, 4)", "(B, 32, 16, 32)", ""),
    ("V = Y W_V", "(4, 1)", "(B, 32, 16, 32)", ""),
    ("S = Q K^T", "(4, 4)", "(B, 32, 16, 16)", "every query against every slot, per head"),
    ("attention", "(4, 4)", "(B, 32, 16, 16)", "each row adds up to 1"),
    ("O = attention V", "(4, 1)", "(B, 32, 16, 32)", "one blend per head per query"),
    ("sharpness", "10 (set by hand)", "1/sqrt(32) x learned", "real scores are scaled by 1/sqrt(d_k)"),
], [("The code: src/hemo/model.py, class CrossAttn",
     src(h_model.CrossAttn.__init__) + [""] + src(h_model.CrossAttn._split) + [""] + src(h_model.CrossAttn.heads))])

side_panel(wb["3 Two heads"], [
    ("heads", "2", "32", ""),
    ("valves g", "(2,)", "(32,), copied to (B, 32)", "one gain per head"),
    ("[g_1 O_1 | g_2 O_2]", "(4, 2)", "(B, 16, 1024)", "32 heads x 32 numbers, side by side"),
    ("W_O", "(2, 2)", "(1024, 64)", ""),
    ("prediction", "(4, 2)", "(B, 16, 64)", ""),
    ("loss", "average of 8", "average of B x 16 x 64", "mean squared error"),
], [("The code: src/hemo/model.py, CrossAttn.combine and forward",
     src(h_model.CrossAttn.combine) + [""] + src(h_model.CrossAttn.forward))])

side_panel(wb["4 Why R heads"], [
    ("unknown items per query", "2", "4", "each 16 numbers long in the real task"),
    ("heads needed", "2", "4", "one per distance"),
    ("mixing matrix", "(2, 2)", "(heads used, 4)", "rows = the heads' recipes"),
    ("re-fit data", "1 query", "16,384 rows", "1024 examples x 16 query tokens"),
    ("re-fit design matrix", "-", "(16,384, 32k + 1)", "k heads x 32 numbers, plus a constant"),
], [("The code: experiments/equations_test.py (the real check of one head, one equation)",
     file_src("experiments/equations_test.py", "def head_outputs", "return float("))])

side_panel(wb["5 Collateral value"], [
    ("heads valued", "3 (A, B, C)", "32", "every 25 training steps"),
    ("head outputs Z", "(4, 3)", "(8,192, 1,024)", "512 probe examples x 16 tokens; 32 heads x 32"),
    ("Gram matrix A = Z^T Z / n", "(3, 3)", "(1,024, 1,024)", "plus a tiny ridge for stability"),
    ("W_O re-fit", "(3, 2)", "(1,024, 64)", "W = M B, with M the inverse of A for open heads"),
    ("value of head h", "re-fit by hand", "trace(W_J^T M_JJ^-1 W_J) / 64", "same number as re-fitting without h, without redoing it"),
], [("The code: src/hemo/model.py, HemoAttn.probe_ischemia",
     src(h_model.HemoAttn.probe_ischemia, "H, dk = self.H", "self.head_value.copy_(value"))])

train_lines = [l for l in src(h_train.train) if any(k in l for k in (
    "price = cfg.price_frac", "for step in range(total)", "model.relax_tone()",
    "if (step - hold) % cfg.probe_every == 0", "model.probe_ischemia(*make_batch", "model.local_step(price)"))]
side_panel(wb["6 The rule"], [
    ("heads at the start", "3", "32", ""),
    ("price", "1 (set by hand)", "0.03 x L_triv", "L_triv = loss of a model that always guesses the average"),
    ("decisions", "2 rounds", "every 25 steps, step 400 to 4000", "one closure (or reopening) per decision"),
    ("fade", "instant", "1/100 per step", "taper = 100"),
    ("reopen", "not shown", "if worth > 2 x price", "so a head at the margin does not flicker"),
], [("The code: src/hemo/model.py, HemoAttn.local_step and relax_tone",
     src(h_model.HemoAttn.local_step) + [""] + src(h_model.HemoAttn.relax_tone)),
    ("The code: src/hemo/train.py, the lines of train() that run the rule",
     list(dict.fromkeys(l.strip() for l in train_lines)))])   # each line once (one-shot repeats two)

side_panel(wb["7 Backup heads"], [
    ("R", "set above", "2 or 4", ""),
    ("failure chance p", "set above", "0.05, 0.1, 0.2, 0.3", "head_dropout"),
    ("price", "set above", "0.03 x L_triv", ""),
    ("failure patterns", "exact formula", "32 random patterns per probe", "probe_masks = 32"),
], [("The code: src/hemo/model.py, HemoAttn._probe_under_damage",
     src(h_model.HemoAttn._probe_under_damage))])

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
os.makedirs(os.path.dirname(OUT), exist_ok=True)
try:
    wb.save(OUT)
except PermissionError:                           # the file is open in Excel: save beside it
    OUT = OUT.replace(".xlsx", " (new).xlsx")
    wb.save(OUT)
    print("hmha_by_hand.xlsx is open in Excel; close it and rerun to replace it")
print("wrote", OUT)
